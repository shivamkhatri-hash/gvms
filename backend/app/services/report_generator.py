# pyright: reportMissingTypeStubs=false
import io
import csv
from typing import List, Dict, Any, cast
import matplotlib  # type: ignore
matplotlib.use('Agg')  # Set non-interactive backend at module level before importing pyplot
import matplotlib.pyplot as plt  # type: ignore
import numpy as np

from openpyxl import Workbook  # type: ignore
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side  # type: ignore
from openpyxl.utils import get_column_letter  # type: ignore
from reportlab.lib.pagesizes import letter, A4, landscape  # type: ignore
from reportlab.lib import colors  # type: ignore
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable, Image  # type: ignore
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle  # type: ignore
from app.models.registry import VariableRegistry
from typing import Optional


class ReportGenerator:
    @staticmethod
    def generate_csv(variables: List[VariableRegistry], records: List[Dict[str, Any]]) -> bytes:
        """Dynamically generate a CSV byte stream using the registered columns."""
        output = io.StringIO()
        writer = csv.writer(output)
        
        # Headers: display_name
        headers = [v.display_name for v in variables]
        writer.writerow(headers)
        
        for r in records:
            row = []
            for v in variables:
                col_name = cast(str, v.sql_column_name)
                val = r.get(col_name)
                row.append(val if val is not None else "")
            writer.writerow(row)
            
        return output.getvalue().encode("utf-8")

    @staticmethod
    def generate_excel(variables: List[VariableRegistry], records: List[Dict[str, Any]], sheet_name: str = "Dataset Report") -> bytes:
        """Dynamically generate an Excel workbook using the registered column definitions and formatting."""
        wb = Workbook()
        ws = wb.active
        if ws is None:
            ws = wb.create_sheet()
        # Excel sheet title sanitization: replace invalid characters \ / ? * : [ ] with _
        cleaned_title = sheet_name
        for char in ['\\', '/', '?', '*', ':', '[', ']']:
            cleaned_title = cleaned_title.replace(char, '_')
        ws.title = cleaned_title[:30]  # Excel limits sheet name to 31 chars

        # Styles
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="003366", end_color="003366", fill_type="solid")  # ONGC Cobalt Blue
        thin_border = Border(
            left=Side(style="thin", color="CCCCCC"),
            right=Side(style="thin", color="CCCCCC"),
            top=Side(style="thin", color="CCCCCC"),
            bottom=Side(style="thin", color="CCCCCC")
        )
        align_center = Alignment(horizontal="center", vertical="center")
        align_left = Alignment(horizontal="left", vertical="center")

        # Headers display_name + unit
        headers = []
        for v in variables:
            h = v.display_name
            if v.display_unit:
                h += f" ({v.display_unit})"
            headers.append(h)
            
        ws.append(headers)

        # Style header row
        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = align_center

        # Append records
        for row_idx, r in enumerate(records, start=2):
            row_data = []
            for v in variables:
                col_name = cast(str, v.sql_column_name)
                val = r.get(col_name)
                if val is None:
                    row_data.append("")
                else:
                    if v.is_numeric:
                        try:
                            row_data.append(float(val))
                        except (ValueError, TypeError):
                            row_data.append(str(val))
                    else:
                        row_data.append(str(val))
            ws.append(row_data)
            
            # Format row cells
            for col_num in range(1, len(headers) + 1):
                cell = ws.cell(row=row_idx, column=col_num)
                cell.border = thin_border
                if variables[col_num - 1].is_numeric:
                    cell.alignment = align_center
                else:
                    cell.alignment = align_left

        # Auto-fit columns
        for col_idx, col in enumerate(ws.columns, start=1):
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col_idx)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 50)

        output = io.BytesIO()
        wb.save(output)
        return output.getvalue()

    @staticmethod
    def generate_pdf(
        dataset_display_name: str,
        variables: List[VariableRegistry],
        records: List[Dict[str, Any]],
        selected_graphs: Optional[List[Dict[str, Any]]] = None
    ) -> bytes:
        """Dynamically generate a styled landscape or portrait PDF containing the dataset and summary."""
        # 1. Filter columns that are visible (max 8 columns for layout rendering)
        pdf_vars = [v for v in variables if v.is_visible][:8]
        if not pdf_vars:
            pdf_vars = variables[:6]

        # 2. Determine page layout: Landscape for wide tables
        use_landscape = len(pdf_vars) > 5
        page_size = landscape(A4) if use_landscape else A4
        printable_width = (842 - 72) if use_landscape else (595 - 72)

        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=page_size,
            rightMargin=36,
            leftMargin=36,
            topMargin=36,
            bottomMargin=36
        )
        elements = []
        styles = getSampleStyleSheet()

        # Custom paragraph styles
        title_style = ParagraphStyle(
            'TitleStyle',
            parent=styles['Heading1'],
            fontName='Helvetica-Bold',
            fontSize=16,
            textColor=colors.HexColor('#003366'),
            spaceAfter=4
        )
        subtitle_style = ParagraphStyle(
            'SubTitleStyle',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=9,
            textColor=colors.HexColor('#64748B'),
            spaceAfter=12
        )
        section_style = ParagraphStyle(
            'SectionStyle',
            parent=styles['Heading2'],
            fontName='Helvetica-Bold',
            fontSize=11,
            textColor=colors.HexColor('#003366'),
            spaceBefore=8,
            spaceAfter=6
        )

        elements.append(Paragraph("OIL AND NATURAL GAS CORPORATION LIMITED", title_style))
        elements.append(Paragraph(f"Dynamic Laboratory Data Platform — {dataset_display_name} Technical Report", subtitle_style))
        elements.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor('#003366'), spaceAfter=12))

        # 3. Dynamic Executive Summary stats
        total = len(records)
        numeric_kpis = [v for v in variables if v.is_numeric and v.kpi_enabled]
        summary_parts = []
        
        for v in numeric_kpis:
            col = cast(str, v.sql_column_name)
            vals = [float(r[col]) for r in records if r.get(col) is not None]
            if vals:
                avg_val = sum(vals) / len(vals)
                summary_parts.append(f"Avg {v.display_name}: <b>{avg_val:.2f} {v.display_unit or ''}</b>")
                
        summary_text = f"<b>Summary Stats:</b> Total Samples: <b>{total}</b>"
        if summary_parts:
            summary_text += " | " + " · ".join(summary_parts)
            
        elements.append(Paragraph(summary_text, styles['Normal']))
        elements.append(Spacer(1, 10))

        # 4. Data Table
        elements.append(Paragraph("Ingested Data & Subsurface Records", section_style))

        # Header row
        table_data = [[
            Paragraph(f"<b>{v.display_name}</b>", styles['Normal'])
            for v in pdf_vars
        ]]

        # Data rows (limit top 100 to avoid out-of-memory or huge page count crashes)
        for r in records[:100]:
            row = []
            for v in pdf_vars:
                col_name = cast(str, v.sql_column_name)
                val = r.get(col_name)
                if val is None:
                    txt = "-"
                else:
                    if v.is_numeric:
                        try:
                            txt = f"{float(val):.2f}"
                        except (ValueError, TypeError):
                            txt = str(val)
                    else:
                        txt = str(val)
                row.append(Paragraph(txt, styles['Normal']))
            table_data.append(row)

        # Distribute width equally
        col_width = printable_width / len(pdf_vars)
        col_widths = [col_width] * len(pdf_vars)

        t = Table(table_data, colWidths=col_widths)
        t.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#F1F5F9')),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ]))
        elements.append(t)

        # 5. Selected Subsurface Charts & Visualizations
        if selected_graphs and records:
            elements.append(Spacer(1, 20))
            elements.append(Paragraph("Visualizations & Dynamic Interpretations", section_style))
            elements.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor('#CBD5E1'), spaceAfter=15))
            
            for g in selected_graphs:
                try:
                    x_col = g.get("x_axis")
                    if not x_col or not isinstance(x_col, str):
                        continue
                    y_col = g.get("y_axis")
                    if y_col is not None and not isinstance(y_col, str):
                        y_col = None
                    title = g.get("title", "Dataset Visualization")
                    g_type = g.get("type", "scatter")
                    
                    # Resolve professional labels from variables registry
                    x_label = x_col.upper()
                    for v in variables:
                        if v.sql_column_name.lower() == x_col.lower():
                            unit = f" ({v.display_unit})" if v.display_unit else ""
                            x_label = f"{v.display_name}{unit}"
                            break
                    
                    y_label = y_col.upper() if y_col else None
                    if y_col:
                        for v in variables:
                            if v.sql_column_name.lower() == y_col.lower():
                                unit = f" ({v.display_unit})" if v.display_unit else ""
                                y_label = f"{v.display_name}{unit}"
                                break

                    print(f"DEBUG_LABEL: x_col={x_col} -> x_label={x_label} | y_col={y_col} -> y_label={y_label}", flush=True)

                    img_bytes = ReportGenerator._create_matplotlib_plot(
                        chart_type=g_type,
                        x_col=x_col,
                        y_col=y_col,
                        title=title,
                        records=records,
                        x_label=x_label,
                        y_label=y_label
                    )
                    
                    if img_bytes:
                        chart_img = Image(io.BytesIO(img_bytes), width=360, height=240)
                        chart_img.hAlign = 'CENTER'
                        elements.append(chart_img)
                        elements.append(Spacer(1, 15))
                except Exception as ex:
                    print(f"Error generating matplotlib plot in PDF: {ex}")
                    elements.append(Paragraph(f"[!] Error rendering visualization chart '{g.get('title')}'", styles['Normal']))
                    elements.append(Spacer(1, 10))

        doc.build(elements)
        buffer.seek(0)
        return buffer.getvalue()

    @staticmethod
    def _create_matplotlib_plot(
        chart_type: str,
        x_col: str,
        y_col: Optional[str],
        title: str,
        records: List[Dict[str, Any]],
        x_label: str,
        y_label: Optional[str]
    ) -> bytes:
        """Dynamically render custom matplotlib charts matching the look & feel of frontend Plotly charts."""
        try:
            # Case-insensitive / underscore-insensitive helper to lookup keys in records
            def get_val_case_insensitive(record: dict, key: str):
                if not key:
                    return None
                if key in record:
                    return record[key]
                key_clean = key.lower().replace("_", "")
                for k, v in record.items():
                    if k.lower().replace("_", "") == key_clean:
                        return v
                return None

            # Extract coordinates
            x_vals = []
            y_vals = []
            
            if chart_type != 'csia_profile':
                for r in records:
                    x_val = get_val_case_insensitive(r, x_col)
                    y_val = get_val_case_insensitive(r, y_col) if y_col else None
                    
                    if x_val is not None:
                        try:
                            x_vals.append(float(x_val))
                        except (ValueError, TypeError):
                            continue
                    else:
                        continue
                        
                    if y_col:
                        if y_val is not None:
                            try:
                                y_vals.append(float(y_val))
                            except (ValueError, TypeError):
                                x_vals.pop()  # Maintain alignment
                                continue
                        else:
                            x_vals.pop()
                            continue

                if not x_vals:
                    return b""

            # Setup figure
            fig, ax = plt.subplots(figsize=(6, 5))
            
            # Apply modern premium styling matching dashboard theme
            ax.set_facecolor('#ffffff')
            fig.patch.set_facecolor('#ffffff')
            ax.spines['top'].set_visible(False)
            ax.spines['right'].set_visible(False)
            ax.spines['left'].set_color('#cbd5e1')
            ax.spines['bottom'].set_color('#cbd5e1')
            ax.spines['left'].set_linewidth(0.8)
            ax.spines['bottom'].set_linewidth(0.8)
            ax.tick_params(colors='#475569', labelsize=8)
            ax.grid(True, linestyle='--', alpha=0.3, color='#cbd5e1', zorder=1)

            # Common helper to plot grouped multi-well data points dynamically
            def plot_grouped_scatter(x_c, y_c):
                groups = {}
                for r in records:
                    x_v = get_val_case_insensitive(r, x_c)
                    y_v = get_val_case_insensitive(r, y_c) if y_c else None
                    if x_v is not None and y_v is not None:
                        try:
                            xf = float(x_v)
                            yf = float(y_v)
                            well = get_val_case_insensitive(r, "well_name") or get_val_case_insensitive(r, "name") or get_val_case_insensitive(r, "well") or "Unknown"
                            if well not in groups:
                                groups[well] = {"x": [], "y": []}
                            groups[well]["x"].append(xf)
                            groups[well]["y"].append(yf)
                        except (ValueError, TypeError):
                            continue

                well_colors = ['#0284c7', '#ec4899', '#dc2626', '#16a34a', '#ea580c', '#d946ef', '#10b981', '#8b5cf6']
                well_markers = ['d', 's', 'o', '^', 'P', '*', 'X', '_']
                
                for idx, (well_name, pts) in enumerate(groups.items()):
                    name_upper = well_name.upper()
                    if name_upper == 'A' or 'WELL A' in name_upper:
                        symbol = 'd'; color = '#0284c7'
                    elif name_upper == 'B' or 'WELL B' in name_upper:
                        symbol = 's'; color = '#ec4899'
                    elif name_upper == 'C' or 'WELL C' in name_upper:
                        symbol = 'o'; color = '#dc2626'
                    elif name_upper == 'D' or 'WELL D' in name_upper:
                        symbol = '^'; color = '#16a34a'
                    elif name_upper == 'E' or 'WELL E' in name_upper:
                        symbol = 'P'; color = '#ea580c'
                    elif name_upper == 'F' or 'WELL F' in name_upper:
                        symbol = '*'; color = '#d946ef'
                    elif name_upper == 'G' or 'WELL G' in name_upper:
                        symbol = 'X'; color = '#10b981'
                    elif name_upper == 'H' or 'WELL H' in name_upper:
                        symbol = '_'; color = '#8b5cf6'
                    else:
                        color = well_colors[idx % len(well_colors)]
                        symbol = well_markers[idx % len(well_markers)]

                    ax.scatter(
                        pts["x"], pts["y"],
                        label=well_name,
                        color=color,
                        marker=symbol,
                        alpha=0.9,
                        edgecolors='black' if symbol not in ['_', '*'] else color,
                        linewidths=0.5,
                        s=35,
                        zorder=5
                    )
                return len(groups) > 0

            # Apply styling matching dashboard charts
            if chart_type == 's2_vs_toc':
                # Group by TOC Classification dynamically to match frontend PlotlyTocS2Scatter colors
                grades = ['Poor', 'Fair', 'Good', 'Very Good', 'Excellent']
                toc_colors = {
                    'Poor': '#EF4444',
                    'Fair': '#F59E0B',
                    'Good': '#10B981',
                    'Very Good': '#06B6D4',
                    'Excellent': '#8B5CF6'
                }
                
                groups = {g: {"x": [], "y": []} for g in grades}
                groups["Other"] = {"x": [], "y": []}
                
                for r in records:
                    x_v = get_val_case_insensitive(r, x_col)
                    y_v = get_val_case_insensitive(r, y_col) if y_col else None
                    if x_v is not None and y_v is not None:
                        try:
                            xf = float(x_v)
                            yf = float(y_v)
                            cls = get_val_case_insensitive(r, "toc_classification") or "Other"
                            if cls not in groups:
                                cls = "Other"
                            groups[cls]["x"].append(xf)
                            groups[cls]["y"].append(yf)
                        except (ValueError, TypeError):
                            continue

                for grade in grades + ["Other"]:
                    pts = groups[grade]
                    if not pts["x"]:
                        continue
                    color = toc_colors.get(grade, '#64748B')
                    ax.scatter(
                        pts["x"], pts["y"],
                        label=f"TOC Grade: {grade}" if grade != "Other" else "Other",
                        color=color,
                        alpha=0.85,
                        edgecolors='white',
                        linewidths=0.5,
                        s=35,
                        zorder=5
                    )
                
                ax.set_xscale('log')
                ax.set_yscale('log')
                ax.set_xlim(0.1, 100)
                ax.set_ylim(0.1, 100)
                
                # Classifications vertical reference lines (TOC guidelines)
                ax.axvline(0.5, color='#06B6D4', linestyle='--', linewidth=0.8, alpha=0.7, zorder=2)
                ax.axvline(1.0, color='#F97316', linestyle='--', linewidth=0.8, alpha=0.7, zorder=2)
                ax.axvline(2.0, color='#2563EB', linestyle='--', linewidth=0.8, alpha=0.7, zorder=2)
                ax.axvline(4.0, color='#EF4444', linestyle='--', linewidth=0.8, alpha=0.7, zorder=2)
                
                # Classifications horizontal reference lines (S2 guidelines)
                ax.axhline(2.5, color='#2563EB', linestyle='--', linewidth=0.8, alpha=0.7, zorder=2)
                ax.axhline(5.0, color='#EF4444', linestyle='--', linewidth=0.8, alpha=0.7, zorder=2)
                ax.axhline(10.0, color='#16A34A', linestyle='--', linewidth=0.8, alpha=0.7, zorder=2)
                ax.axhline(20.0, color='#7C3AED', linestyle='--', linewidth=0.8, alpha=0.7, zorder=2)
                
                ax.legend(loc='lower left', fontsize=7, frameon=True, facecolor='#ffffff', edgecolor='#e2e8f0')

            elif chart_type == 'hi_vs_tmax':
                plot_grouped_scatter(x_col, y_col)
                ax.set_xlim(400, 480)
                ax.set_ylim(0, 700)
                
                # Thermal Maturity stage vertical division boundaries
                ax.axvline(435, color='#64748B', linestyle='--', linewidth=0.8, alpha=0.7, zorder=2)
                ax.axvline(470, color='#64748B', linestyle='--', linewidth=0.8, alpha=0.7, zorder=2)
                
                # Zone Annotation Labels
                ax.text(417.5, 665, "Immature", fontsize=6, color='#64748B', fontweight='bold', ha='center', va='center')
                ax.text(452.5, 665, "Mature", fontsize=6, color='#64748B', fontweight='bold', ha='center', va='center')
                ax.text(475.0, 665, "Post-Mature", fontsize=6, color='#64748B', fontweight='bold', ha='center', va='center')

                # Maturity curve guidelines (Type II, III, etc.)
                tmax_vals = np.linspace(400, 472, 100)
                y_type_iii = 85 * (1 - 0.7 * ((tmax_vals - 400)/72)**2)
                y_type_ii = 630 * (1 - 0.9 * ((tmax_vals - 400)/72)**2)
                
                ax.plot(tmax_vals, y_type_iii, color='#F97316', linestyle='-', linewidth=1, label='Type III', zorder=2)
                ax.plot(tmax_vals, y_type_ii, color='#1E3A8A', linestyle='-', linewidth=1, label='Type II', zorder=2)
                ax.legend(loc='upper right', fontsize=7, frameon=True, facecolor='#ffffff', edgecolor='#e2e8f0')
                
            elif chart_type == 'depth_profile':
                plot_grouped_scatter(x_col, y_col)
                ax.invert_yaxis()  # Invert depth log
                ax.legend(loc='upper right', fontsize=7, frameon=True, facecolor='#ffffff', edgecolor='#e2e8f0')

            elif chart_type == 'api_vs_depth':
                # Group by well name dynamically
                groups = {}
                for r in records:
                    x_v = get_val_case_insensitive(r, x_col)
                    y_v = get_val_case_insensitive(r, y_col) if y_col else None
                    if x_v is not None and y_v is not None:
                        try:
                            xf = float(x_v)
                            yf = float(y_v)
                            well = get_val_case_insensitive(r, "well_name") or get_val_case_insensitive(r, "name") or get_val_case_insensitive(r, "well") or "Unknown"
                            if well not in groups:
                                groups[well] = {"x": [], "y": []}
                            groups[well]["x"].append(xf)
                            groups[well]["y"].append(yf)
                        except (ValueError, TypeError):
                            continue
                
                well_colors = ['#2563EB', '#DC2626', '#16A34A', '#7C3AED', '#F59E0B', '#EC4899', '#10B981', '#6366F1']
                for idx, (well_name, pts) in enumerate(groups.items()):
                    color = well_colors[idx % len(well_colors)]
                    ax.scatter(pts["x"], pts["y"], label=well_name, color=color, alpha=0.8, edgecolors='black', linewidths=0.5, s=35, zorder=5)
                
                ax.set_xlim(10, 70)
                ax.set_ylim(4000, 1000) # Inverted
                
                # Dashed vertical classification boundaries
                ax.axvline(30, color='#000000', linestyle='--', linewidth=0.8, alpha=0.5)
                ax.axvline(40, color='#000000', linestyle='--', linewidth=0.8, alpha=0.5)
                ax.axvline(60, color='#000000', linestyle='--', linewidth=0.8, alpha=0.5)
                
                # Text labels matching frontend:
                ax.text(20, 1200, "Heavy oils", fontsize=7, fontweight='bold', color='#000000', ha='center', va='center')
                ax.text(35, 1200, "Medium oils", fontsize=7, fontweight='bold', color='#000000', ha='center', va='center')
                ax.text(50, 1200, "Light oils", fontsize=7, fontweight='bold', color='#000000', ha='center', va='center')
                ax.text(65, 1200, "Condensates /\nvery light oils", fontsize=7, fontweight='bold', color='#000000', ha='center', va='center')
                
                x_label = "API (°API)"
                y_label = "Depth (m)"
                if groups:
                    ax.legend(loc='lower left', fontsize=7)

            elif chart_type == 'pr_nc17_vs_ph_nc18':
                groups = {}
                for r in records:
                    x_v = get_val_case_insensitive(r, x_col)
                    y_v = get_val_case_insensitive(r, y_col) if y_col else None
                    if x_v is not None and y_v is not None:
                        try:
                            xf = float(x_v)
                            yf = float(y_v)
                            well = get_val_case_insensitive(r, "name") or get_val_case_insensitive(r, "well_name") or get_val_case_insensitive(r, "well") or "Unknown"
                            if well not in groups:
                                groups[well] = {"x": [], "y": []}
                            groups[well]["x"].append(xf)
                            groups[well]["y"].append(yf)
                        except (ValueError, TypeError):
                            continue

                # Map markers & colors to match Well A to H style
                for well_name, pts in groups.items():
                    name_upper = well_name.upper()
                    symbol = 'o'
                    color = '#3b82f6'
                    
                    if name_upper == 'A' or 'WELL A' in name_upper or 'WELL_A' in name_upper:
                        symbol = 'd'
                        color = '#0284c7'
                    elif name_upper == 'B' or 'WELL B' in name_upper or 'WELL_B' in name_upper:
                        symbol = 's'
                        color = '#ec4899'
                    elif name_upper == 'C' or 'WELL C' in name_upper or 'WELL_C' in name_upper:
                        symbol = 'o'
                        color = '#dc2626'
                    elif name_upper == 'D' or 'WELL D' in name_upper or 'WELL_D' in name_upper:
                        symbol = '^'
                        color = '#16a34a'
                    elif name_upper == 'E' or 'WELL E' in name_upper or 'WELL_E' in name_upper:
                        symbol = 'P'
                        color = '#ea580c'
                    elif name_upper == 'F' or 'WELL F' in name_upper or 'WELL_F' in name_upper:
                        symbol = '*'
                        color = '#d946ef'
                    elif name_upper == 'G' or 'WELL G' in name_upper or 'WELL_G' in name_upper:
                        symbol = 'X'
                        color = '#10b981'
                    elif name_upper == 'H' or 'WELL H' in name_upper or 'WELL_H' in name_upper:
                        symbol = '_'
                        color = '#8b5cf6'
                    else:
                        hash_val = sum(ord(char) for char in name_upper)
                        colors_list = ['#0891b2', '#0d9488', '#4f46e5', '#7c3aed', '#db2777', '#ca8a04']
                        symbols_list = ['o', '^', 'v', 'd', 's', 'x']
                        color = colors_list[hash_val % len(colors_list)]
                        symbol = symbols_list[hash_val % len(symbols_list)]
                    
                    ax.scatter(
                        pts["x"], pts["y"],
                        label=well_name,
                        color=color,
                        marker=symbol,
                        alpha=0.9,
                        edgecolors='black' if symbol not in ['_', '*'] else color,
                        linewidths=0.5,
                        s=35,
                        zorder=5
                    )
                
                ax.set_xscale('log')
                ax.set_yscale('log')
                ax.set_xlim(0.01, 10.0)
                ax.set_ylim(0.01, 10.0)
                
                # Constant Pr/Ph ratios:
                x_line = np.logspace(-2, 1, 100)
                ax.plot(x_line, 8.0 * x_line, color='#000000', linewidth=0.6, zorder=2)
                ax.plot(x_line, 4.0 * x_line, color='#000000', linewidth=0.6, zorder=2)
                ax.plot(x_line, 2.0 * x_line, color='#000000', linewidth=0.6, zorder=2)
                ax.plot(x_line, 1.0 * x_line, color='#000000', linewidth=0.6, zorder=2)
                ax.plot(x_line, 0.5 * x_line, color='#000000', linewidth=0.6, zorder=2)

                # Annotation labels parallel to diagonals
                ax.text(0.80, 6.4, "Terrestrial, Type III", fontsize=5.5, color='#000000', rotation=35, ha='center', va='center')
                ax.text(1.20, 4.8, "Terrestrial, CoalyType III", fontsize=5.5, color='#000000', rotation=35, ha='center', va='center')
                ax.text(1.50, 3.0, "Type II-Type III mixture", fontsize=5.5, color='#000000', rotation=35, ha='center', va='center')
                ax.text(2.20, 2.2, "Type II, reducing algal, marine", fontsize=5.5, color='#000000', rotation=35, ha='center', va='center')

                # Biodegradation text and arrow
                ax.text(0.12, 1.2, "Biodegradation", fontsize=6, fontweight='bold', color='#000000', rotation=35, ha='center', va='center')
                ax.annotate("", xy=(0.20, 1.8), xytext=(0.08, 0.70), arrowprops=dict(arrowstyle="->", color="black", lw=0.8))

                # Maturation text and arrow
                ax.text(0.15, 0.05, "Maturation", fontsize=6, fontweight='bold', color='#000000', rotation=35, ha='center', va='center')
                ax.annotate("", xy=(0.09, 0.03), xytext=(0.22, 0.070), arrowprops=dict(arrowstyle="->", color="black", lw=0.8))

                # Oxidizing & Reducing double-headed arrow and text
                ax.annotate("", xy=(0.52, 0.85), xytext=(0.75, 0.45), arrowprops=dict(arrowstyle="<->", color="black", lw=0.8))
                ax.text(0.55, 0.92, "Oxidizing", fontsize=6, fontweight='bold', color='#000000', rotation=-55, ha='center', va='center')
                ax.text(0.70, 0.38, "Reducing", fontsize=6, fontweight='bold', color='#000000', rotation=-55, ha='center', va='center')

                x_label = r'Phytane / $\mathrm{nC}_{18}$'
                y_label = r'Pristane / $\mathrm{nC}_{17}$'
                if groups:
                    ax.legend(loc='lower left', fontsize=7)
                
            elif chart_type == 'sofer_plot' or chart_type == 'galimov_plot':
                plot_grouped_scatter(x_col, y_col)
                # Plot the Sofer boundary line (CV = 0.47)
                # CV = -2.53 * sat + 2.22 * aro - 11.65
                # For CV = 0.47: aro = (2.53 * sat + 12.12) / 2.22
                sat_line = np.linspace(-35, -20, 100)
                aro_line = (2.53 * sat_line + 12.12) / 2.22
                ax.plot(sat_line, aro_line, color='#dc2626', linestyle='--', linewidth=1, label='Sofer Line (CV=0.47)', zorder=2)
                ax.text(-25, -23.5, "Terrigenous (waxy)", fontsize=7, color='#dc2626', fontweight='bold', ha='center')
                ax.text(-32, -31, "Marine (non-waxy)", fontsize=7, color='#0284c7', fontweight='bold', ha='center')
                ax.set_xlim(-35, -20)
                ax.set_ylim(-35, -20)
                ax.legend(loc='lower right', fontsize=7)
                x_label = "δ13C Saturates (‰)"
                y_label = "δ13C Aromatics (‰)"

            elif chart_type == 'csia_profile':
                well_colors = ['#0284c7', '#ec4899', '#dc2626', '#16a34a', '#ea580c', '#d946ef', '#10b981', '#8b5cf6']
                well_markers = ['d', 's', 'o', '^', 'P', '*', 'X', '_']
                
                carbons = list(range(15, 35))
                for idx, r in enumerate(records[:15]):  # Limit to top 15 to avoid clutter
                    well_name = get_val_case_insensitive(r, "well_name") or get_val_case_insensitive(r, "name") or "Unknown"
                    depth = get_val_case_insensitive(r, "interval_top") or get_val_case_insensitive(r, "depth") or "N/A"
                    
                    x_lbls = []
                    y_vals = []
                    for c_num in carbons:
                        v = get_val_case_insensitive(r, f"nc{c_num}")
                        if v is not None:
                            try:
                                y_vals.append(float(v))
                                x_lbls.append(f"nC{c_num}")
                            except (ValueError, TypeError):
                                pass
                                
                    if y_vals:
                        color = well_colors[idx % len(well_colors)]
                        marker = well_markers[idx % len(well_markers)]
                        ax.plot(x_lbls, y_vals, label=f"{well_name} ({depth}m)", color=color, marker=marker, markersize=4, linewidth=1, alpha=0.9)
                
                x_label = "n-Alkanes"
                y_label = "δ13C (‰)"
                ax.legend(loc='upper right', fontsize=6, frameon=True, facecolor='#ffffff', edgecolor='#e2e8f0')

            else: # Standard scatter
                ax.scatter(x_vals, y_vals, color='#10B981', alpha=0.8, edgecolors='white', linewidths=0.5, s=35, zorder=5)

            # Automatically invert Y-axis if label/column suggests depth profiles
            if y_col and any(d in y_col.lower() for d in ['depth', 'top', 'bottom', 'interval_top']):
                if not ax.yaxis_inverted():
                    ax.invert_yaxis()

            # Log-scale Bernard Diagram isotope ratios automatically
            if x_col and x_col.lower() == 'c1_by_c2_plus_c3':
                ax.set_xscale('log')
                ax.set_xlim(1.0, 10000.0)

            ax.set_title(title, fontsize=10, fontweight='bold', color='#0F172A', pad=8)
            ax.set_xlabel(x_label, fontsize=8, color='#475569')
            ax.set_ylabel(y_label or 'Value', fontsize=8, color='#475569')
            
            # Hide grid for pr_nc17_vs_ph_nc18 to match reference
            if chart_type == 'pr_nc17_vs_ph_nc18':
                ax.grid(False)
            else:
                ax.grid(True, which='major', linestyle='--', color='#CBD5E1', linewidth=0.5, alpha=0.4)
            
            for spine in ['top', 'right']:
                ax.spines[spine].set_visible(False)
            for spine in ['left', 'bottom']:
                ax.spines[spine].set_color('#CBD5E1')
                ax.spines[spine].set_linewidth(0.8)

            buf = io.BytesIO()
            plt.savefig(buf, format='png', bbox_inches='tight', dpi=150)
            plt.close(fig)
            return buf.getvalue()
        except Exception as e:
            print(f"Failed to render matplotlib plot: {e}")
            return b""
